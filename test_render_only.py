"""Renders all formats from sample data and checks each file is valid. No Ollama needed."""
import json, zipfile, py_compile, csv, xml.etree.ElementTree as ET
from html.parser import HTMLParser
from schema import Document, Section, Table
from renderers import render_all

doc = Document(
    title="AI-Based Attendance System", subtitle="Project Report",
    sections=[
        Section(heading="Introduction",
                paragraphs=["Attendance tracking is slow when done manually. This project uses face recognition."],
                bullets=["Contactless", "Real-time logging", "Reduces proxy attendance"]),
        Section(heading="Accuracy Results", paragraphs=["Tested on three classrooms."],
                table=Table(headers=["Classroom", "Students", "Accuracy (%)"],
                            rows=[["Room A", "40", "96"], ["Room B", "55", "93"], ["Room C", "30", "98"]])),
        Section(heading="Conclusion", paragraphs=["Works well in normal lighting."], bullets=["Low cost"]),
    ])

paths = render_all(doc, base="sample")


def check(fmt, p):
    if fmt == "json": json.load(open(p, encoding="utf-8"))
    elif fmt == "xml": ET.parse(p)
    elif fmt == "py": py_compile.compile(p, doraise=True)
    elif fmt == "ipynb": assert json.load(open(p, encoding="utf-8"))["nbformat"] == 4
    elif fmt == "csv": assert len(list(csv.reader(open(p, encoding="utf-8-sig")))) > 5
    elif fmt == "html": HTMLParser().feed(open(p, encoding="utf-8").read())
    elif fmt == "doc":
        s = open(p, encoding="ascii").read()
        assert s.startswith("{\\rtf1") and s.count("{") - s.count("\\{") == s.count("}") - s.count("\\}")
    elif fmt == "zip": assert zipfile.ZipFile(p).testzip() is None
    elif fmt == "odt":
        from odf.opendocument import load; load(p)
    elif fmt == "docx":
        import docx; docx.Document(p)
    elif fmt == "pptx":
        from pptx import Presentation; Presentation(p)
    elif fmt == "xlsx":
        from openpyxl import load_workbook; load_workbook(p)
    elif fmt == "pdf": assert open(p, "rb").read(4) == b"%PDF"


for fmt, p in paths.items():
    check(fmt, p)
    print(f"OK  {fmt.upper():6} {p}")
print(f"\nAll {len(paths)} formats valid.")

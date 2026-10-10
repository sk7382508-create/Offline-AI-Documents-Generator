import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from schema import Document, Section

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0xE0, 0x3A, 0x3E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x22, 0x22)


def _bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def _text(slide, x, y, w, h, text, size, color, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tf


def _points(sec: Section, limit=5):
    """Max 5 short bullet points per slide."""
    pts = list(sec.bullets)
    if not pts:
        for p in sec.paragraphs:
            pts += [s.strip() for s in re.split(r"(?<=[.!?])\s+", p) if s.strip()][:2]
    return pts[:limit]


def render(doc: Document, path: str):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # title slide
    s = prs.slides.add_slide(blank)
    _bg(s, NAVY)
    _text(s, 1, 2.6, 11.3, 1.5, doc.title, 44, WHITE, True)
    _text(s, 1, 4.3, 11.3, 1, doc.subtitle, 22, RGBColor(0xCC, 0xD6, 0xE5))

    for i, sec in enumerate(doc.sections, 1):
        s = prs.slides.add_slide(blank)
        _bg(s, WHITE)
        bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.3))
        bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
        _text(s, 0.6, 0.3, 12, 0.8, f"{i}. {sec.heading}", 30, WHITE, True)
        tf = _text(s, 0.8, 1.7, 11.7, 5, "", 20, DARK)
        first = True
        for pt in _points(sec):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            r = p.add_run(); r.text = "•  " + pt
            r.font.size = Pt(20); r.font.color.rgb = DARK
            p.space_after = Pt(12)

        if sec.table and sec.table.headers:
            t = prs.slides.add_slide(blank)
            _bg(t, WHITE)
            _text(t, 0.6, 0.3, 12, 0.8, f"{sec.heading}: data", 28, NAVY, True)
            n = len(sec.table.headers)
            rows = sec.table.rows[:10]
            shape = t.shapes.add_table(len(rows) + 1, n, Inches(0.8), Inches(1.5),
                                       Inches(11.7), Inches(0.5 * (len(rows) + 1)))
            tbl = shape.table
            for c, h in enumerate(sec.table.headers):
                cell = tbl.cell(0, c); cell.text = h
                cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = WHITE
            for ri, row in enumerate(rows, 1):
                row = (list(row) + [""] * n)[:n]
                for c, v in enumerate(row):
                    cell = tbl.cell(ri, c); cell.text = str(v)
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(12)

    # closing slide
    s = prs.slides.add_slide(blank)
    _bg(s, NAVY)
    _text(s, 1, 3, 11.3, 1.5, "Thank You", 48, WHITE, True)
    prs.save(path)

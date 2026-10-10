import re
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from schema import Document

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def _num(v):
    s = str(v).strip().replace(",", "")
    for cast in (int, float):
        try:
            return cast(s)
        except ValueError:
            pass
    return v


def _sheet_name(name, used):
    base = re.sub(r"[\[\]:*?/\\]", "", name)[:28] or "Sheet"
    n, k = base, 2
    while n in used:
        n = f"{base[:25]}_{k}"; k += 1
    used.add(n)
    return n


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autowidth(ws, cap=60):
    for col in ws.columns:
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(w + 3, cap)


def render(doc: Document, path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    used = {"Overview"}

    ws["A1"] = doc.title; ws["A1"].font = Font(bold=True, size=16, color="1F3A5F")
    ws["A2"] = doc.subtitle; ws["A2"].font = Font(italic=True)
    ws.append([])
    ws.append(["#", "Section", "Summary", "Key Points"])
    _style_header(ws, 4, 4)
    for i, sec in enumerate(doc.sections, 1):
        ws.append([i, sec.heading,
                   sec.paragraphs[0] if sec.paragraphs else "",
                   "\n".join("• " + b for b in sec.bullets)])
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 60

    for sec in doc.sections:
        if not (sec.table and sec.table.headers and sec.table.rows):
            continue
        n = len(sec.table.headers)
        t = wb.create_sheet(_sheet_name(sec.heading, used))
        t.append(sec.table.headers)
        _style_header(t, 1, n)
        rows = [[_num(v) for v in (list(r) + [""] * n)[:n]] for r in sec.table.rows]
        for r in rows:
            t.append(r)
        _autowidth(t)

        # native chart on the first fully numeric column
        for c in range(1, n):
            if all(isinstance(r[c], (int, float)) for r in rows):
                ch = BarChart()
                ch.title = sec.heading
                ch.add_data(Reference(t, min_col=c + 1, min_row=1, max_row=len(rows) + 1),
                            titles_from_data=True)
                ch.set_categories(Reference(t, min_col=1, min_row=2, max_row=len(rows) + 1))
                ch.width, ch.height = 16, 8
                t.add_chart(ch, f"{chr(ord('A') + n + 1)}2")
                break
    wb.save(path)

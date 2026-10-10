from docx import Document as Docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from schema import Document


def render(doc: Document, path: str):
    d = Docx()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    # cover page
    for _ in range(6):
        d.add_paragraph()
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run(doc.title)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    s = d.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.add_run(doc.subtitle).font.size = Pt(14)
    d.add_page_break()

    # sections
    for i, sec in enumerate(doc.sections, 1):
        d.add_heading(f"{i}. {sec.heading}", level=1)
        for p in sec.paragraphs:
            d.add_paragraph(p)
        for b in sec.bullets:
            d.add_paragraph(b, style="List Bullet")
        if sec.table and sec.table.headers:
            tbl = d.add_table(rows=1, cols=len(sec.table.headers))
            tbl.style = "Light Grid Accent 1"
            for c, h in enumerate(sec.table.headers):
                tbl.rows[0].cells[c].text = h
            for row in sec.table.rows:
                cells = tbl.add_row().cells
                for c in range(len(sec.table.headers)):
                    cells[c].text = row[c] if c < len(row) else ""
    d.save(path)
